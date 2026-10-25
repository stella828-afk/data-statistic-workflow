"""读取“抖音订单”，计算原表“周赛”里的抖音业绩。"""

from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


# 第一步：设置输入文件和输出文件。以后换文件，先修改这里。
SOURCE_FILE = Path(r"D:\【个人工作】\mamiya实习\A#周业绩播报\播报（周）7-10.xlsx")
OUTPUT_FILE = Path(__file__).with_name("周赛_抖音结果.xlsx")

# 左边是“周赛”的昵称，右边是订单表“归属达人运营”的姓名。
# 按原表上方计算区的顺序输出，不重新排名。
BD_NAMES = {
    "召召": "田明召",
    "C哥": "陈希",
    "佳美": "吴佳美",
    "鸭鸭": "鸭鸭",
    "玉磊": "王玉磊",
    "佳妮": "骆佳妮",
    "石头": "陈石",
}


def calculate_totals(source_file):
    """第二步：按姓名汇总计佣金额，相当于原表的 SUMIF。"""
    totals = {name: Decimal("0") for name in BD_NAMES.values()}
    row_count = other_count = ignored_count = 0
    book = load_workbook(source_file, read_only=True, data_only=True)
    try:
        if "抖音订单" not in book.sheetnames:
            raise ValueError("输入文件里没有“抖音订单”工作表。")
        rows = book["抖音订单"].iter_rows(values_only=True)
        headers = next(rows, ())
        required = ("归属达人运营", "计佣金额")
        if any(headers.count(name) != 1 for name in required):
            raise ValueError("第一行必须各有一列“归属达人运营”和“计佣金额”。")
        owner_col = headers.index("归属达人运营")
        amount_col = headers.index("计佣金额")

        for row_number, row in enumerate(rows, start=2):
            if all(value is None for value in row):
                continue
            row_count += 1
            owner, amount = row[owner_col], row[amount_col]
            if owner not in totals:
                other_count += 1
                continue
            if isinstance(amount, str) and amount.startswith("#"):
                raise ValueError(f"第 {row_number} 行计佣金额存在 Excel 错误：{amount}")
            # SUMIF 对金额列中的空白、文本和布尔值不做数值累加。
            if not isinstance(amount, (int, float)) or isinstance(amount, bool):
                ignored_count += 1
                continue
            money = Decimal(str(amount))  # 用十进制计算，避免小数累加误差。
            if not money.is_finite():
                raise ValueError(f"第 {row_number} 行的计佣金额不是有限数值。")
            totals[owner] += money
    finally:
        book.close()
    return totals, row_count, other_count, ignored_count


def save_results(totals, output_file):
    """第三步：另存结果，原始文件保持不变。"""
    book = Workbook()
    sheet = book.active
    sheet.title = "周赛"
    sheet.append(["BD", "抖音"])
    for nickname, name in BD_NAMES.items():
        amount = totals[name].quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        sheet.append([nickname, float(amount)])
    # 以下仅调整结果表的字体、列宽和金额格式，不影响计算。
    for row in sheet:
        for cell in row:
            cell.font = Font(name="微软雅黑", size=11)
            cell.alignment = Alignment(vertical="center")
    for cell in sheet[1]:
        cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="23465C")
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "#,##0.00"
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 20
    for row_number in range(1, 9):
        sheet.row_dimensions[row_number].height = 24
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    book.save(output_file)
    book.close()


def main():
    if SOURCE_FILE.resolve() == OUTPUT_FILE.resolve():
        raise ValueError("输出路径不能和原始输入文件相同。")
    totals, rows, other, ignored = calculate_totals(SOURCE_FILE)
    save_results(totals, OUTPUT_FILE)
    print(f"已读取 {rows} 行订单。以下金额单位为元：")
    for nickname, name in BD_NAMES.items():
        print(f"  {nickname}：{totals[name]:,.2f}")
    print(f"未列入这 7 位 BD 的订单：{other} 行；匹配行中忽略的非数字金额：{ignored} 行。")
    print(f"结果已保存：{OUTPUT_FILE}")


if __name__ == "__main__":
    try:
        main()
    except (OSError, ValueError) as error:
        print(f"运行未完成：{error}")
        print("请检查输入路径；如果结果文件已在 Excel 中打开，请先关闭再运行。")
        sys.exit(1)
